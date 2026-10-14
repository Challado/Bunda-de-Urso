"""Coleta: CKAN com descoberta de recurso, páginas oficiais, CSV/XLSX/JSON.
Falha de API/HTML ou mudança de layout -> falha explícita; nunca inventar dados.
"""
import csv, io, json, re, logging
from urllib.parse import urljoin, urlparse
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from defusedxml import ElementTree as ET
from . import settings
from .normalization import key
log=logging.getLogger(__name__)

class SourceError(RuntimeError):pass
class NotModified(SourceError):pass

class Downloader:
    """HTTPS, limite de download, redirects verificados e retry limitado.
    Uma falha externa nunca elimina o cache anteriormente aprovado.
    """
    def __init__(self, session=None, sleep=None, attempts=3):
        import time
        self.session=session or requests.Session()
        self.sleep=sleep or time.sleep
        self.attempts=attempts
        self.session.headers.update({
            'User-Agent':settings.USER_AGENT,
            'Accept':'application/json,text/csv,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,text/html,application/xml,application/octet-stream;q=0.8'})

    @staticmethod
    def checked_url(url):
        import ipaddress
        u=urlparse(url)
        if u.scheme!='https' or not u.hostname or u.username or u.password:
            raise SourceError('Download exige URL HTTPS pública sem credenciais embutidas')
        host=u.hostname.lower()
        if host in ('localhost','localhost.localdomain') or host.endswith(('.localhost','.local','.internal')):
            raise SourceError('Host interno não permitido na coleta')
        try:
            ipaddress.ip_address(host)
            raise SourceError('Endereço IP literal não permitido na coleta')
        except ValueError:pass
        if u.port not in (None,443):raise SourceError('Porta externa não permitida')
        return url

    def fetch(self,url,etag=None,last_modified=None,method='GET',data=None):
        import time, email.utils
        from datetime import datetime, timezone
        self.checked_url(url)
        headers={}
        if etag:headers['If-None-Match']=etag
        if last_modified:headers['If-Modified-Since']=last_modified
        last_error=None
        for attempt in range(self.attempts):
            target=url
            try:
                for hop in range(6):
                    self.checked_url(target)
                    with self.session.request(method,target,data=data,headers=headers,
                        timeout=(settings.CONNECT_TIMEOUT,settings.REQUEST_TIMEOUT),stream=True,allow_redirects=False) as res:
                        if res.status_code in (301,302,303,307,308):
                            location=res.headers.get('Location')
                            if not location:raise SourceError('Redirect sem Location')
                            target=urljoin(target,location)
                            self.checked_url(target)
                            if res.status_code==303:method='GET';data=None
                            continue
                        if res.status_code==304:raise NotModified('HTTP 304')
                        if res.status_code in (429,500,502,503,504):
                            retry_after=res.headers.get('Retry-After','')
                            retry_seconds=None
                            try:retry_seconds=max(0,min(30,int(retry_after)))
                            except ValueError:
                                try:
                                    date=email.utils.parsedate_to_datetime(retry_after)
                                    retry_seconds=max(0,min(30,(date-datetime.now(timezone.utc)).total_seconds()))
                                except (ValueError,TypeError,OverflowError):pass
                            raise RetriableSourceError(f'HTTP {res.status_code} em {target}',retry_seconds)
                        if res.status_code>=400:raise SourceError(f'HTTP {res.status_code} em {target}')
                        size=0;blocks=[]
                        for chunk in res.iter_content(chunk_size=65536):
                            size+=len(chunk)
                            if size>settings.MAX_BYTES:raise SourceError('Download acima de RADAR_MAX_BYTES')
                            if chunk:blocks.append(chunk)
                        raw=b''.join(blocks)
                        if not raw:raise SourceError('HTTP 200 com corpo vazio')
                        return raw,dict(res.headers),target
                raise SourceError('Mais de cinco redirecionamentos HTTP')
            except NotModified:raise
            except (requests.exceptions.RequestException,TimeoutError,RetriableSourceError) as e:
                last_error=e
                if attempt==self.attempts-1:break
                delay=e.wait if isinstance(e,RetriableSourceError) and e.wait is not None else min(15,2**attempt)
                log.warning('Tentativa %s/%s falhou (%s); retry em %ss',attempt+1,self.attempts,e,delay)
                self.sleep(delay)
        raise SourceError(f'HTTP falhou após {self.attempts} tentativas: {last_error}') from last_error

class RetriableSourceError(SourceError):
    def __init__(self,message,wait=None):
        super().__init__(message);self.wait=wait

def ckan_discover(source,downloader):
    """Preferência por API CKAN; fallback HTML. Escolhe UM arquivo mais recente."""
    preference=[p.upper() for p in source.get('resource_formats',['JSON','CSV','XLSX'])]
    resources=[]
    if source.get('ckan_api'):
        try:
            body,_,_=downloader.fetch(source['ckan_api'])
            obj=json.loads(body.decode('utf-8-sig'))
            if not obj.get('success') or not isinstance(obj.get('result',{}).get('resources'),list):raise SourceError('CKAN sem resources')
            for r in obj['result']['resources']:
                fmt=(r.get('format') or '').upper();url=r.get('url')
                if url and fmt in preference and urlparse(url).scheme=='https':
                    resources.append({'url':url,'format':fmt,'updated':r.get('last_modified') or r.get('created') or '', 'name':r.get('name') or ''})
        except (SourceError,ValueError,KeyError) as e:
            log.warning('CKAN API inacessível; tentando página HTML: %s',e)
    if not resources:
        body,_,url=downloader.fetch(source['url']);soup=BeautifulSoup(body,'html.parser')
        for a in soup.select('a[href]'):
            link=urljoin(url,a['href']);txt=' '.join(a.get_text(' ',strip=True).split())
            ext=re.search(r'\.(xlsx|csv|json)(?:\?|$)',link,re.I)
            fmt=ext.group(1).upper() if ext else None
            if not fmt:
                ctx=a.find_parent(class_=re.compile('resource-item|resource',re.I))
                cl=(ctx.get_text(' ',strip=True) if ctx else txt).lower()
                for f in preference:
                    if f.lower() in cl:fmt=f;break
            if fmt in preference and urlparse(link).scheme=='https' and ('download' in link or ext):
                resources.append({'url':link,'format':fmt,'updated':'','name':txt})
    if not resources:raise SourceError('Sem recursos válidos CSV/JSON/XLSX; revisar layout/licenciamento')
    # CKAN tipicamente retorna recursos da versão mais recente por created/last_modified
    # priorizar data de nome quando disponível; empates por formato preferido
    def rank(r):
        mm=re.search(r'(20\d\d)[-/_. ]([01]?\d)',r['name']+' '+r['url'])
        period=(int(mm.group(1)),int(mm.group(2))) if mm and 1<=int(mm.group(2))<=12 else (0,0)
        return (period,r['updated'], -preference.index(r['format']))
    winner=max(resources,key=rank)
    return winner

def prf_fixed_discover(source,downloader):
    """Segue link oficial PRF e encontra XLSX exportável na tabela remota.
    Não usar o HTML dinâmico como se contivesse registros renderizados.
    """
    body,_,url=downloader.fetch(source['url'])
    soup=BeautifulSoup(body,'html.parser')
    table_url=None
    for a in soup.select('a[href]'):
        if 'TABELA DE RADARES FIXOS' in a.get_text(' ',strip=True).upper():
            table_url=urljoin(url,a['href']);break
    if not table_url:raise SourceError('PRF: link oficial TABELA DE RADARES FIXOS ausente')
    if urlparse(table_url).hostname not in ('mapas-html.vercel.app',):
        raise SourceError('PRF: destino de tabela mudou; revisão de origem necessária')
    html,_,url=downloader.fetch(table_url)
    inner=BeautifulSoup(html,'html.parser')
    for a in inner.select('a[href]'):
        link=urljoin(url,a['href'])
        u=urlparse(link)
        if u.hostname=='docs.google.com' and '/spreadsheets/d/' in u.path and ('export?' in link or 'format=xlsx' in link):
            return {'url':link,'format':'xlsx','name':'Planilha publicada via link TABELA DE RADARES FIXOS da PRF'}
    raise SourceError('PRF: planilha XLSX não encontrada no HTML dinâmico; conferir link oficial')

def parse_html(raw,fmt):
    soup=BeautifulSoup(raw,'html.parser')
    if fmt=='prf_portateis':
        rows=[]
        for tr in soup.select('tr'):
            tds=[x.get_text(' ',strip=True) for x in tr.select('td')]
            if len(tds)>=2 and re.fullmatch(r'[A-Z]{2,5}\d{2,4}',tds[0].replace(' ',''),re.I):
                rows.append({'identificador':tds[0],'numero_serie':tds[1],'tipo':'portatil'})
        if not rows:
            text=soup.get_text('\n')
            for line in text.splitlines():
                m=re.match(r'^\s*([A-Z]{2,5}\d{2,4})\s*[|\t ]+([A-Z0-9\-]{3,})\s*$',line,re.I)
                if m: rows.append({'identificador':m.group(1),'numero_serie':m.group(2),'tipo':'portatil'})
        if not rows:raise SourceError('Layout PRF portáteis não reconhecido, zero registros')
        return rows
    tables=soup.select('table')
    for table in tables:
        headings=[x.get_text(' ',strip=True) for x in table.select('thead th')]
        if not headings:
            first=table.select_one('tr')
            if first:headings=[x.get_text(' ',strip=True) for x in first.select('th')]
        if len(headings)<3:continue
        if not any('radar' in key(x) or 'identificador' in key(x) or 'rodovia' in key(x) for x in headings):continue
        result=[]
        for tr in table.select('tbody tr') or table.select('tr')[1:]:
            values=[td.get_text(' ',strip=True) for td in tr.select('td')]
            if len(values)>=len(headings):result.append(dict(zip(headings,values)))
        if result:return result
    raise SourceError('Tabela HTML não disponível como dados server-side; fonte requer adaptador de layout')

def _feature_collection(raw):
    """Importação de pontos GeoJSON; não converter segmentos ou polígonos em radares.

    Coordenadas GeoJSON são sempre [longitude, latitude]. Preserva propriedades e
    geometria de origem em cada registro para auditoria de qualidade.
    """
    try:
        obj=json.loads(raw.decode('utf-8-sig'))
    except (UnicodeError, ValueError) as e:
        raise SourceError('GeoJSON inválido') from e
    if not isinstance(obj,dict) or obj.get('type')!='FeatureCollection' or not isinstance(obj.get('features'),list):
        raise SourceError('GeoJSON deve ser FeatureCollection com features')
    rows=[]
    for i,feature in enumerate(obj['features']):
        if not isinstance(feature,dict) or feature.get('type')!='Feature':
            raise SourceError(f'Feature GeoJSON inválida na posição {i}')
        geom=feature.get('geometry')
        if not isinstance(geom,dict) or geom.get('type')!='Point':
            # Segmentos de trecho elegível não são radares pontuais.
            continue
        coordinates=geom.get('coordinates')
        if not isinstance(coordinates,(list,tuple)) or len(coordinates)<2:
            raise SourceError(f'Coordenadas GeoJSON ausentes na posição {i}')
        try:
            lon,lat=float(coordinates[0]),float(coordinates[1])
        except (TypeError,ValueError,OverflowError) as e:
            raise SourceError(f'Coordenadas GeoJSON inválidas na posição {i}') from e
        import math
        if not (math.isfinite(lat) and math.isfinite(lon) and -90<=lat<=90 and -180<=lon<=180):
            raise SourceError(f'Coordenadas GeoJSON fora da faixa na posição {i}')
        props=feature.get('properties') or {}
        if not isinstance(props,dict):raise SourceError(f'properties inválido na posição {i}')
        for k in ('latitude','longitude'):
            if k in props and str(props[k]).strip() not in ('',str(lat) if k=='latitude' else str(lon)):
                from .normalization import decimal
                expected=lat if k=='latitude' else lon
                normalized=decimal(props[k])
                if normalized is not None and abs(float(normalized)-expected)>0.000001:
                    raise SourceError(f'Conflito entre {k} e geometria GeoJSON na posição {i}')
        row=dict(props);row['latitude']=lat;row['longitude']=lon
        if feature.get('id') is not None and 'identificador' not in row:
            row['identificador']=str(feature['id'])
        row['_geometria_original']=geom
        rows.append(row)
    if not rows:
        raise SourceError('GeoJSON sem pontos de radar; feições lineares não podem virar pontos')
    return rows


def _kml_points(raw):
    """Lê KML Placemarks pontuais sem transformar áreas/linhas em equipamentos."""
    try:
        root=ET.fromstring(raw)
    except Exception as e:raise SourceError('KML/XML inválido') from e
    rows=[]
    def children_text(node,name):
        return next((el.text.strip() for el in node.iter()
                     if el.tag.split('}')[-1]==name and el.text and el.text.strip()),None)
    for item in root.iter():
        if item.tag.split('}')[-1]!='Placemark':continue
        points=[x for x in item.iter() if x.tag.split('}')[-1]=='Point']
        if not points:continue
        if len(points)!=1:raise SourceError('Placemark KML com múltiplos pontos ambíguos')
        coords=children_text(points[0],'coordinates')
        if not coords:raise SourceError('Placemark KML sem coordinates')
        values=coords.strip().split(',')
        if len(values)<2:raise SourceError('Coordenadas KML inválidas')
        try:lon,lat=float(values[0]),float(values[1])
        except (TypeError,ValueError,OverflowError) as e:raise SourceError('Coordenadas KML inválidas') from e
        import math
        if not (math.isfinite(lat) and math.isfinite(lon) and -90<=lat<=90 and -180<=lon<=180):
            raise SourceError('Coordenadas KML fora da faixa')
        row={'latitude':lat,'longitude':lon}
        if children_text(item,'name'):row['identificador']=children_text(item,'name')
        if children_text(item,'description'):row['observacao_origem']=children_text(item,'description')
        for element in item.iter():
            if element.tag.split('}')[-1]=='Data' and element.get('name'):
                v=children_text(element,'value')
                if v is not None:row[element.get('name')]=v
            if element.tag.split('}')[-1]=='SimpleData' and element.get('name') and element.text:
                row[element.get('name')]=element.text.strip()
        rows.append(row)
    if not rows:raise SourceError('KML sem Placemarks com Point válidos')
    return rows


def parse_table(raw,fmt):
    fmt=fmt.lower()
    if fmt=='geojson':return _feature_collection(raw)
    if fmt=='kml':return _kml_points(raw)
    if fmt in ('prf_portateis','prf_fixos','html'):return parse_html(raw,fmt)
    if fmt=='xlsx':
        if not raw.startswith(b'PK'):raise SourceError('XLSX sem assinatura ZIP')
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            best=[]
            for sheet in wb:
                iterator=sheet.iter_rows(values_only=True)
                for i,line in enumerate(iterator):
                    headers=[str(x).strip() if x is not None else '' for x in line]
                    if 2<=len([x for x in headers if x]) and any(key(x) in ('uf','rodovia','latitude','identificador','identificadordoequipamento','municipio','br','km','codigoequipamento','concessionaria') for x in headers):
                        rows=[]
                        for r in iterator:
                            if any(x is not None for x in r):rows.append(dict(zip(headers,[str(x).strip() if x is not None else None for x in r])))
                        if len(rows)>len(best):best=rows
                        break
                    if i>=30:break
            if not best:raise SourceError('Nenhuma aba XLSX com cabeçalho de radar reconhecido')
            return best
        finally:wb.close()
    if fmt=='csv':
        for encoding in ('utf-8-sig','cp1252','latin-1'):
            try:decoded=raw.decode(encoding);break
            except UnicodeError:continue
        if decoded.lstrip().startswith('<'):raise SourceError('HTML recebido no lugar de CSV')
        try:dialect=csv.Sniffer().sniff(decoded[:4096],delimiters=',;\t|')
        except csv.Error:dialect=csv.excel
        result=list(csv.DictReader(io.StringIO(decoded),dialect=dialect))
        if not result or len(result[0])<2:raise SourceError('CSV sem linhas/cabeçalho válido')
        return result
    if fmt in ('json','manual_json'):
        try:data=json.loads(raw.decode('utf-8-sig'))
        except (ValueError,UnicodeError) as e:raise SourceError('JSON inválido') from e
        if isinstance(data,list):rows=data
        elif isinstance(data,dict):
            rows=next((data[k] for k in ('result','records','data','radares','medidores') if isinstance(data.get(k),list)),None)
            if rows is None:raise SourceError('JSON sem array de registros reconhecido')
        else:raise SourceError('JSON inesperado')
        if not rows or not all(isinstance(r,dict) for r in rows):raise SourceError('JSON sem registros válidos')
        return rows
    if fmt=='xml':
        try:root=ET.fromstring(raw)
        except Exception as e:raise SourceError('XML inválido') from e
        def unpack(node):
            # Conservar estruturas e atributos. Não achatar Faixas/Histórico.
            children=list(node)
            if not children:
                text=(node.text or '').strip()
                if node.attrib:return {'_atributos':dict(node.attrib),'_texto':text or None}
                return text or None
            pairs={}
            for child in children:
                k=child.tag.split('}')[-1]
                pairs.setdefault(k,[]).append(unpack(child))
            obj={k:(v if len(v)>1 else v[0]) for k,v in pairs.items()}
            if node.attrib:obj['_atributos']=dict(node.attrib)
            return obj
        rows=[]
        for node in root:
            value=unpack(node)
            if isinstance(value,dict):rows.append(value)
            elif isinstance(value,list):rows.extend(x for x in value if isinstance(x,dict))
        if not rows:raise SourceError('XML sem registros estruturados reconhecíveis')
        return rows
    raise SourceError('Formato não suportado: '+fmt)

OVERPASS_QUERY='[out:json][timeout:60];(nwr["highway"="speed_camera"](-34.0,-74.0,5.4,-34.0);nwr["enforcement"="maxspeed"](-34.0,-74.0,5.4,-34.0););out center tags;'

def osm_parse(raw):
    try:data=json.loads(raw)
    except ValueError as e:raise SourceError('Overpass inválido') from e
    out=[]
    for e in data.get('elements',[]):
        tags=e.get('tags') or {}
        lat=e.get('lat') or (e.get('center') or {}).get('lat');lon=e.get('lon') or (e.get('center') or {}).get('lon')
        if lat is None or lon is None:continue
        out.append({'identificador':f"OSM-{e.get('type')}-{e.get('id')}",'latitude':lat,'longitude':lon,'tipo':'fixo','rodovia':tags.get('ref') or tags.get('name'),'velocidade':tags.get('maxspeed'),'sentido':tags.get('direction')})
    if not out:raise SourceError('OSM sem radares/elementos')
    return out
